import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"D:\lyx\Downloads\deepseek\report.md"
OUT = r"D:\lyx\Downloads\deepseek\report.docx"

doc = Document()

# base style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def clean(text):
    # strip markdown links to "text (url)" and bold markers lightly
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "")
    return text.strip()

def add_code_line(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    return p

lines = open(SRC, encoding="utf-8").read().splitlines()

i = 0
in_code = False
while i < len(lines):
    line = lines[i]
    if line.strip().startswith("```"):
        in_code = not in_code
        i += 1
        continue
    if in_code:
        add_code_line(line)
        i += 1
        continue
    s = line.strip()
    if not s:
        i += 1
        continue
    if s.startswith("### "):
        doc.add_heading(clean(s[4:]), level=3)
    elif s.startswith("## "):
        doc.add_heading(clean(s[3:]), level=2)
    elif s.startswith("# "):
        doc.add_heading(clean(s[2:]), level=1)
    elif s.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run(clean(s[2:]))
        run.italic = True
    elif s.startswith("- ") or s.startswith("* "):
        doc.add_paragraph(clean(s[2:]), style="List Bullet")
    elif s.startswith("|"):
        # table block
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            rows.append(cells)
            i += 1
        if len(rows) >= 2:
            header = rows[0]
            data = rows[2:] if rows[1] and set("".join(rows[1])).issubset("-: ") else rows[1:]
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            hdr = t.rows[0].cells
            for j, h in enumerate(header):
                hdr[j].text = clean(h)
            for r in data:
                cells = t.add_row().cells
                for j, c in enumerate(r):
                    if j < len(cells):
                        cells[j].text = clean(c)
        continue
    else:
        doc.add_paragraph(clean(s))
    i += 1

doc.save(OUT)
print("saved", OUT)
